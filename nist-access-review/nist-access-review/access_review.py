#!/usr/bin/env python3
"""NIST 800-53 user access review automation.

Reads a user-access export (CSV), runs the rules in rules.yaml, maps every finding to
NIST SP 800-53 controls, and writes a reviewer-ready Excel workbook plus an HTML summary.

Usage:
    python access_review.py --input data/sample_users.csv --rules rules.yaml --out output
"""
from __future__ import annotations

import argparse
import html
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REQUIRED_COLUMNS = [
    "username", "department", "job_title", "manager", "account_status", "last_login",
    "created_date", "groups", "privileged", "privileged_justified_date", "account_type",
    "termination_date", "expiry_date", "interactive_login",
]
DATE_COLUMNS = ["last_login", "created_date", "termination_date", "expiry_date",
                "privileged_justified_date"]


# --------------------------------------------------------------------------- loading
def load_users(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise SystemExit(f"Input is missing required columns: {', '.join(missing)}")
    for col in df.columns:
        df[col] = df[col].str.strip()
    for col in DATE_COLUMNS:
        df[col] = pd.to_datetime(df[col], errors="coerce")
    for col in ("privileged", "interactive_login", "account_type", "account_status"):
        df[col] = df[col].str.lower()
    return df


# --------------------------------------------------------------------------- checks
# Each check returns a list of (row, detail_text). Add a new check = add a function here
# and register it in CHECKS, then reference it from rules.yaml.
def check_inactive_enabled(df, p, as_of):
    limit, skip = p.get("days", 90), set(p.get("exclude_types", []))
    out = []
    for _, r in df.iterrows():
        if r["account_status"] != "enabled" or r["account_type"] in skip:
            continue
        if pd.isna(r["last_login"]):
            age = (as_of - r["created_date"]).days if pd.notna(r["created_date"]) else None
            if age is not None and age > limit:
                out.append((r, f"No login on record; account is {age} days old"))
        else:
            idle = (as_of - r["last_login"]).days
            if idle > limit:
                out.append((r, f"Last login {r['last_login']:%Y-%m-%d} ({idle} days ago)"))
    return out


def check_terminated_active(df, p, as_of):
    out = []
    for _, r in df.iterrows():
        t = r["termination_date"]
        if pd.notna(t) and t <= as_of and r["account_status"] == "enabled":
            out.append((r, f"Terminated {t:%Y-%m-%d} ({(as_of - t).days} days ago); account still enabled"))
    return out


def check_contractor_expired(df, p, as_of):
    out = []
    for _, r in df.iterrows():
        if r["account_type"] != "contractor" or r["account_status"] != "enabled":
            continue
        e = r["expiry_date"]
        if pd.isna(e):
            out.append((r, "Contractor account has no expiry date set"))
        elif e < as_of:
            out.append((r, f"Expired {e:%Y-%m-%d} ({(as_of - e).days} days ago); still enabled"))
    return out


def check_privileged_unjustified(df, p, as_of):
    limit, out = p.get("max_age_days", 180), []
    for _, r in df.iterrows():
        if r["privileged"] != "yes" or r["account_status"] != "enabled":
            continue
        j = r["privileged_justified_date"]
        if pd.isna(j):
            out.append((r, "Privileged access with no justification on record"))
        elif (as_of - j).days > limit:
            out.append((r, f"Privilege last justified {j:%Y-%m-%d} ({(as_of - j).days} days ago; limit {limit})"))
    return out


def check_sod_conflict(df, p, as_of):
    out = []
    for _, r in df.iterrows():
        groups = {g.strip() for g in r["groups"].split(";") if g.strip()}
        for a, b in p.get("conflicts", []):
            if a in groups and b in groups:
                out.append((r, f"Holds conflicting groups: {a} + {b}"))
    return out


def check_no_manager(df, p, as_of):
    return [(r, "Manager/owner field is empty")
            for _, r in df.iterrows()
            if r["account_status"] == "enabled" and not r["manager"]]


def check_service_interactive(df, p, as_of):
    return [(r, "Service account has interactive login enabled")
            for _, r in df.iterrows()
            if r["account_type"] == "service" and r["interactive_login"] == "yes"
            and r["account_status"] == "enabled"]


def check_shared_account(df, p, as_of):
    pats = [re.compile(x, re.I) for x in p.get("name_patterns", [])]
    out = []
    for _, r in df.iterrows():
        if r["account_status"] != "enabled":
            continue
        if r["account_type"] == "shared":
            out.append((r, "Account type is 'shared'"))
        elif any(p_.search(r["username"]) for p_ in pats):
            out.append((r, f"Username '{r['username']}' matches a generic-account pattern"))
    return out


CHECKS = {
    "inactive_enabled": check_inactive_enabled,
    "terminated_active": check_terminated_active,
    "contractor_expired": check_contractor_expired,
    "privileged_unjustified": check_privileged_unjustified,
    "sod_conflict": check_sod_conflict,
    "no_manager": check_no_manager,
    "service_interactive": check_service_interactive,
    "shared_account": check_shared_account,
}


# --------------------------------------------------------------------------- engine
def run_review(df, config, as_of):
    weights = config["severity_weights"]
    findings = []
    for rule in config["rules"]:
        fn = CHECKS.get(rule["check"])
        if fn is None:
            raise SystemExit(f"Rule {rule['id']}: unknown check '{rule['check']}'")
        for row, detail in fn(df, rule.get("params") or {}, as_of):
            findings.append({
                "Username": row["username"], "Department": row["department"],
                "Job Title": row["job_title"], "Rule ID": rule["id"], "Issue": rule["name"],
                "Detail": detail, "Severity": rule["severity"],
                "NIST Control": ", ".join(rule["nist"]),
                "Recommended Action": rule["recommendation"],
            })
    findings.sort(key=lambda f: (-weights[f["Severity"]], f["Username"], f["Rule ID"]))
    for i, f in enumerate(findings, 1):
        f["Finding ID"] = f"F-{i:04d}"
    return findings


def risk_by_account(findings, weights):
    acct = defaultdict(lambda: {"score": 0, "issues": [], "dept": ""})
    for f in findings:
        a = acct[f["Username"]]
        a["score"] += weights[f["Severity"]]
        a["issues"].append(f["Issue"])
        a["dept"] = f["Department"]
    return sorted(acct.items(), key=lambda kv: (-kv[1]["score"], kv[0]))


def count_by_family(findings):
    fam = defaultdict(int)
    for f in findings:
        for fam_code in {c.strip().split("-")[0] for c in f["NIST Control"].split(",")}:
            fam[fam_code] += 1
    return dict(sorted(fam.items(), key=lambda kv: -kv[1]))


# --------------------------------------------------------------------------- Excel
FONT = "Arial"
SEV_FILL = {"Critical": ("C00000", "FFFFFF"), "High": ("F4B183", "000000"),
            "Medium": ("FFE699", "000000"), "Low": ("C6E0B4", "000000")}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_excel(findings, config, n_users, as_of, path):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_f = wb.create_sheet("Findings")
    ws_s = wb.create_sheet("Reviewer Sign-off")

    # ---- Findings
    cols = ["Finding ID", "Username", "Department", "Job Title", "Rule ID", "Issue", "Detail",
            "Severity", "NIST Control", "Recommended Action"]
    ws_f.append(cols)
    style_header(ws_f, 1, len(cols))
    for f in findings:
        ws_f.append([f[c] for c in cols])
    last = max(len(findings) + 1, 2)
    for row in ws_f.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        sev = row[7].value
        if sev in SEV_FILL:
            bg, fg = SEV_FILL[sev]
            row[7].fill = PatternFill("solid", fgColor=bg)
            row[7].font = Font(name=FONT, size=10, bold=True, color=fg)
    set_widths(ws_f, [10, 16, 12, 22, 8, 36, 48, 10, 14, 60])
    ws_f.freeze_panes = "A2"
    ws_f.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # ---- Summary (formulas read from the Findings sheet so it recalculates if rows change)
    rng = lambda col: f"Findings!${col}$2:${col}${last}"
    ws_sum["A1"] = "User Access Review - Summary"
    ws_sum["A1"].font = Font(name=FONT, size=14, bold=True)
    kpis = [
        ("As-of date", as_of.isoformat(), None),
        ("Accounts reviewed", n_users, "Count of rows in the input CSV (hardcoded by the script)"),
        ("Total findings", f"=COUNTA({rng('A')})", None),
        ("Accounts with at least one finding", f"=IFERROR(SUMPRODUCT(1/COUNTIF({rng('B')},{rng('B')})),0)", None),
        ("% of accounts with findings", "=IFERROR(B6/B4,0)", None),
    ]
    for i, (label, val, note) in enumerate(kpis, start=3):
        ws_sum.cell(row=i, column=1, value=label).font = Font(name=FONT, bold=True)
        c = ws_sum.cell(row=i, column=2, value=val)
        c.font = Font(name=FONT)
        c.alignment = Alignment(horizontal="right")
        if note:
            ws_sum.cell(row=i, column=3, value=note).font = Font(name=FONT, italic=True, size=9, color="7F7F7F")
    ws_sum["B7"].number_format = "0.0%"

    ws_sum["A9"] = "Findings by severity"
    ws_sum["A9"].font = Font(name=FONT, bold=True, size=12)
    ws_sum.append([])
    ws_sum["A10"], ws_sum["B10"] = "Severity", "Count"
    style_header(ws_sum, 10, 2)
    for i, sev in enumerate(["Critical", "High", "Medium", "Low"], start=11):
        ws_sum.cell(row=i, column=1, value=sev)
        ws_sum.cell(row=i, column=2, value=f'=COUNTIF({rng("H")},A{i})')
        bg, fg = SEV_FILL[sev]
        ws_sum.cell(row=i, column=1).fill = PatternFill("solid", fgColor=bg)
        ws_sum.cell(row=i, column=1).font = Font(name=FONT, bold=True, color=fg)
        ws_sum.cell(row=i, column=2).font = Font(name=FONT)

    ws_sum["A16"] = "Findings by rule and NIST control"
    ws_sum["A16"].font = Font(name=FONT, bold=True, size=12)
    for c, h in enumerate(["Rule ID", "Check", "Severity", "NIST Control", "Count"], 1):
        ws_sum.cell(row=17, column=c, value=h)
    style_header(ws_sum, 17, 5)
    for i, rule in enumerate(config["rules"], start=18):
        vals = [rule["id"], rule["name"], rule["severity"], ", ".join(rule["nist"]),
                f'=COUNTIF({rng("E")},A{i})']
        for c, v in enumerate(vals, 1):
            cell = ws_sum.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT)
            cell.border = BORDER
    first_rule, last_rule = 18, 17 + len(config["rules"])
    chart = BarChart()
    chart.type, chart.title = "bar", "Findings by rule"
    chart.add_data(Reference(ws_sum, min_col=5, min_row=17, max_row=last_rule), titles_from_data=True)
    chart.set_categories(Reference(ws_sum, min_col=1, min_row=first_rule, max_row=last_rule))
    chart.legend, chart.height, chart.width = None, 8, 14
    ws_sum.add_chart(chart, "G9")
    set_widths(ws_sum, [34, 46, 12, 18, 10])
    ws_sum.column_dimensions["B"].width = 46

    # ---- Reviewer sign-off
    ws_s["A1"] = "Reviewer Sign-off"
    ws_s["A1"].font = Font(name=FONT, size=14, bold=True)
    ws_s["A2"] = "Yellow cells are for the reviewer. Decision options: Approve (keep access), Revoke (remove access), Exception (keep with documented approval)."
    ws_s["A2"].font = Font(name=FONT, italic=True, size=9)
    scols = ["Finding ID", "Username", "Issue", "Severity", "NIST Control", "Recommended Action",
             "Decision", "Reviewer", "Decision Date", "Comments"]
    for c, h in enumerate(scols, 1):
        ws_s.cell(row=4, column=c, value=h)
    style_header(ws_s, 4, len(scols))
    for i, f in enumerate(findings, start=5):
        vals = [f["Finding ID"], f["Username"], f["Issue"], f["Severity"], f["NIST Control"],
                f["Recommended Action"], None, None, None, None]
        for c, v in enumerate(vals, 1):
            cell = ws_s.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if c >= 7:
                cell.fill = INPUT_FILL
        ws_s.cell(row=i, column=9).number_format = "yyyy-mm-dd"
    end = max(len(findings) + 4, 5)
    dv = DataValidation(type="list", formula1='"Approve,Revoke,Exception"', allow_blank=True)
    ws_s.add_data_validation(dv)
    dv.add(f"G5:G{end}")
    set_widths(ws_s, [10, 16, 36, 10, 14, 50, 12, 16, 14, 30])
    ws_s.freeze_panes = "A5"
    ws_s.auto_filter.ref = f"A4:J{end}"

    wb.save(path)


# --------------------------------------------------------------------------- HTML
def write_html(findings, config, n_users, as_of, path):
    weights = config["severity_weights"]
    sev_counts = {s: sum(1 for f in findings if f["Severity"] == s) for s in weights}
    accounts = risk_by_account(findings, weights)
    fams = count_by_family(findings)
    rule_counts = {r["id"]: sum(1 for f in findings if f["Rule ID"] == r["id"]) for r in config["rules"]}
    esc = html.escape

    # inline SVG horizontal bar chart (no external libraries)
    mx = max(fams.values(), default=1)
    bars = ""
    for i, (fam, n) in enumerate(fams.items()):
        y, w = 10 + i * 36, int(300 * n / mx)
        bars += (f'<text x="0" y="{y + 17}" class="lbl">{esc(fam)}</text>'
                 f'<rect x="40" y="{y}" width="{w}" height="24" rx="3" class="bar"/>'
                 f'<text x="{46 + w}" y="{y + 17}" class="lbl">{n}</text>')
    chart_h = 20 + 36 * max(len(fams), 1)
    fam_names = {"AC": "Access Control", "IA": "Identification &amp; Authentication", "PS": "Personnel Security"}
    legend = " &middot; ".join(f"{k} = {fam_names.get(k, k)}" for k in fams)

    cards = [("Accounts reviewed", n_users), ("Total findings", len(findings)),
             ("Accounts flagged", len(accounts)),
             ("Critical + High", sev_counts.get("Critical", 0) + sev_counts.get("High", 0))]
    cards_html = "".join(f'<div class="card"><div class="n">{v}</div><div>{esc(k)}</div></div>' for k, v in cards)

    top_rows = "".join(
        f"<tr><td>{i}</td><td><b>{esc(u)}</b></td><td>{esc(d['dept'])}</td><td>{d['score']}</td>"
        f"<td>{esc('; '.join(d['issues']))}</td></tr>"
        for i, (u, d) in enumerate(accounts[:10], 1))
    rule_rows = "".join(
        f"<tr><td>{esc(r['id'])}</td><td>{esc(r['name'])}</td>"
        f"<td><span class='sev {r['severity']}'>{r['severity']}</span></td>"
        f"<td>{esc(', '.join(r['nist']))}</td><td>{rule_counts[r['id']]}</td></tr>"
        for r in config["rules"])

    page = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>User Access Review - {as_of.isoformat()}</title>
<style>
:root{{--bg:#fff;--fg:#1a1a1a;--muted:#666;--card:#f3f5f9;--line:#d9dde5;--bar:#2f5fb3}}
@media (prefers-color-scheme:dark){{:root{{--bg:#16181d;--fg:#e8e8e8;--muted:#9aa0aa;--card:#22252c;--line:#363a44;--bar:#6b9cf0}}}}
body{{font-family:Arial,Helvetica,sans-serif;background:var(--bg);color:var(--fg);max-width:960px;margin:0 auto;padding:24px}}
h1{{margin-bottom:0}} .sub{{color:var(--muted);margin-top:4px}}
.cards{{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}}
.card{{flex:1;min-width:150px;background:var(--card);padding:14px 16px;border-radius:8px}}
.card .n{{font-size:30px;font-weight:bold}}
table{{border-collapse:collapse;width:100%;font-size:14px}}
th,td{{text-align:left;padding:8px;border-bottom:1px solid var(--line);vertical-align:top}}
th{{background:var(--card)}} .wrap{{overflow-x:auto}}
.bar{{fill:var(--bar)}} .lbl{{fill:var(--fg);font-size:14px}}
.sev{{padding:2px 8px;border-radius:10px;font-size:12px;font-weight:bold;color:#000}}
.Critical{{background:#e06666;color:#fff}} .High{{background:#f4b183}} .Medium{{background:#ffe699}} .Low{{background:#c6e0b4}}
.note{{color:var(--muted);font-size:12px;margin-top:24px}}
</style></head><body>
<h1>User Access Review</h1>
<p class="sub">As of {as_of.isoformat()} &middot; findings mapped to NIST SP 800-53 Rev. 5</p>
<div class="cards">{cards_html}</div>
<h2>Findings by NIST control family</h2>
<svg viewBox="0 0 400 {chart_h}" width="100%" style="max-width:480px" role="img" aria-label="Findings by control family">{bars}</svg>
<p class="sub">{legend}</p>
<h2>Top 10 riskiest accounts</h2>
<div class="wrap"><table><tr><th>#</th><th>Account</th><th>Dept</th><th>Risk score</th><th>Issues</th></tr>{top_rows}</table></div>
<h2>Findings by rule</h2>
<div class="wrap"><table><tr><th>ID</th><th>Check</th><th>Severity</th><th>NIST</th><th>Count</th></tr>{rule_rows}</table></div>
<p class="note">Risk score = sum of severity weights ({esc(', '.join(f'{k} {v}' for k, v in weights.items()))}).
Full detail and reviewer sign-off are in findings.xlsx. Sample data is synthetic.</p>
</body></html>"""
    Path(path).write_text(page, encoding="utf-8")


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default="data/sample_users.csv")
    ap.add_argument("--rules", default="rules.yaml")
    ap.add_argument("--out", default="output")
    ap.add_argument("--as-of", default=date.today().isoformat(), help="review date, YYYY-MM-DD")
    args = ap.parse_args()

    as_of = pd.Timestamp(args.as_of)
    config = yaml.safe_load(Path(args.rules).read_text())
    df = load_users(Path(args.input))
    findings = run_review(df, config, as_of)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    write_excel(findings, config, len(df), as_of.date(), out / "findings.xlsx")
    write_html(findings, config, len(df), as_of.date(), out / "report.html")

    print(f"Reviewed {len(df)} accounts -> {len(findings)} findings "
          f"across {len({f['Username'] for f in findings})} accounts")
    for rule in config["rules"]:
        n = sum(1 for f in findings if f["Rule ID"] == rule["id"])
        print(f"  {rule['id']}  {n:3d}  [{rule['severity']:<8}] {rule['name']}")
    print(f"Wrote {out / 'findings.xlsx'} and {out / 'report.html'}")


if __name__ == "__main__":
    main()
